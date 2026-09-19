# -*- coding: utf-8 -*-
"""将「方法流程图」pipeline.png 嵌入论文 4.1 系统总体架构（该章节当前无图）。

插入位置：4.1 标题之后的正文段落后（带图题）。
幂等保护：检测到已嵌入过的图题则跳过。
"""
import os
import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DOCX_PATH = os.path.join(PROJECT_ROOT, "paper", "学年论文.docx")
FIG_PATH = os.path.join(PROJECT_ROOT, "paper", "figures", "pipeline.png")

CAPTION = "图 4.1 基于扩散模型的时间序列异常检测总体架构"


def main():
    if not os.path.exists(FIG_PATH):
        raise SystemExit("流程图不存在: " + FIG_PATH)

    doc = docx.Document(DOCX_PATH)

    # 幂等
    for p in doc.paragraphs:
        if p.text.strip() == CAPTION:
            print("已嵌入过 4.1 流程图，跳过。")
            return

    # 定位 4.1 标题
    target_idx = None
    for i, p in enumerate(doc.paragraphs):
        if p.style.name == "Heading 2" and p.text.startswith("4.1"):
            target_idx = i
            break
    if target_idx is None:
        raise SystemExit("找不到 4.1 标题。")

    # 在 4.1 标题正下方插入图（找到下一个 Heading 之前的位置）
    # 简化：紧跟 4.1 标题段落的下一个正文段落之后
    insert_after = doc.paragraphs[target_idx]
    # 跳过 4.1 标题后的第一个空段，插到第一个含内容的正文段之后
    j = target_idx + 1
    while j < len(doc.paragraphs) and not doc.paragraphs[j].text.strip():
        j += 1  # 跳过空段
    if j < len(doc.paragraphs) and doc.paragraphs[j].style.name != "Heading 2":
        insert_after = doc.paragraphs[j]
    insert_after_el = insert_after._p

    # 构造图与图题（先 append 到文档末尾，再 move 到目标位置）
    fig_para = doc.add_paragraph()
    fig_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fig_para.add_run().add_picture(FIG_PATH, width=Inches(6.5))

    cap_para = doc.add_paragraph()
    cap_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap_run = cap_para.add_run(CAPTION)
    cap_run.bold = True
    cap_run.font.size = Pt(9)

    # 先追加一个空段作为目标（便于两个新段落连续插入）
    anchor_para = doc.add_paragraph("")
    anchor_el = anchor_para._p

    # 把 anchor 移到 insert_after 之后
    insert_after_el.addnext(anchor_el)
    # 把 figure + caption 移到 anchor 之后（依次 addnext）
    anchor_el.addnext(cap_para._p)
    anchor_el.addnext(fig_para._p)
    # 删除临时 anchor
    anchor_el.getparent().remove(anchor_el)

    doc.save(DOCX_PATH)
    print("已嵌入流程图到 4.1 节:", CAPTION)


if __name__ == "__main__":
    main()
